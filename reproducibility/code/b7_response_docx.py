"""Render the point-by-point response letter as a Word document.

Journals often want the response in an editable format, and the LaTeX source is
the only authoritative copy: it carries the generated macros, so a DOCX typed by
hand would drift from the computed numbers the moment anything is recomputed.
This converter expands the same macro files the PDF uses, so the two versions
always report identical values.

    python b7_response_docx.py

Writes Revision_R1_new/response.docx.
"""

import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ANALYSIS = os.path.dirname(HERE)
ROOT = os.path.dirname(os.path.dirname(ANALYSIS))
PKG = os.path.join(ROOT, "revision", "Revision_R1_new")
GEN = os.path.join(PKG, "generated")

REPLY_BLUE = RGBColor(0x17, 0x4B, 0x86)

# ------------------------------------------------------------------- macros
MACRO_FILES = [
    os.path.join(GEN, "results_macros.tex"),
    os.path.join(GEN, "crossrefs.tex"),
    os.path.join(PKG, "si_numbers.tex"),
    os.path.join(PKG, "repository.tex"),
]
DEF = re.compile(r"\\(?:re)?newcommand\*?\s*\{?\\([A-Za-z]+)\}?\s*\{(.*)\}\s*(?:%.*)?$")


def load_macros():
    macros = {}
    for path in MACRO_FILES:
        if not os.path.exists(path):
            continue
        for line in open(path, encoding="utf-8"):
            m = DEF.match(line.strip())
            if m:
                macros[m.group(1)] = m.group(2)
    # the shorthands defined in the document preamble
    macros.update({
        "pic": "pIC50", "rsq": "R2", "qsq": "Q2CV", "dpph": "DPPH*",
        "etal": "et al.", "ie": "i.e.", "eg": "e.g.",
    })
    return macros


MACROS = load_macros()

# --------------------------------------------------------------- inline TeX
SYMBOL = [
    (r"\\rsq\{\}", "R\u00b2"),
    (r"\\qsq\{\}", "Q\u00b2\u1d04\u1d20"),
    (r"\\pic\{\}", "pIC\u2085\u2080"),
    (r"\\dpph\{\}", "DPPH\u2022"),
    (r"\\etal\{\}", "et al."),
    (r"\\ldots\{?\}?", "\u2026"),
    (r"\\dots\{?\}?", "\u2026"),
    (r"\\times", "\u00d7"),
    (r"\\rho", "\u03c1"),
    (r"\\eta", "\u03b7"),
    (r"\\mu", "\u03bc"),
    (r"\\Delta", "\u0394"),
    (r"\\alpha", "\u03b1"),
    (r"\\leq", "\u2264"),
    (r"\\geq", "\u2265"),
    (r"\\pm", "\u00b1"),
    (r"\\approx", "\u2248"),
    (r"\\to", "\u2192"),
    (r"\\rightarrow", "\u2192"),
    (r"\\emph\b", ""),
    (r"\\quad", "    "),
    (r"\\,", " "),
    (r"\\;", " "),
    (r"\\ ", " "),
]

SUPERS = {"2": "\u00b2", "3": "\u00b3"}
# LaTeX accent forms such as \"o are not alphabetic commands, so the generic
# command stripper in detex would leave a stray backslash behind
ACCENT_UML = {"o": "\u00f6", "a": "\u00e4", "u": "\u00fc", "e": "\u00eb", "i": "\u00ef",
              "O": "\u00d6", "A": "\u00c4", "U": "\u00dc"}
ACCENT_ACUTE = {"e": "\u00e9", "a": "\u00e1", "i": "\u00ed", "o": "\u00f3", "u": "\u00fa",
                "c": "\u0107", "n": "\u0144"}
ACCENT_GRAVE = {"e": "\u00e8", "a": "\u00e0", "i": "\u00ec", "o": "\u00f2", "u": "\u00f9"}


def expand_macros(s, depth=0):
    """Substitute \\Name{} tokens with their generated values."""
    if depth > 4:
        return s

    def sub(m):
        return MACROS.get(m.group(1), m.group(0))

    out = re.sub(r"\\([A-Za-z]+)\{\}", sub, s)
    return expand_macros(out, depth + 1) if out != s else out


def detex(s):
    """LaTeX fragment -> readable plain text."""
    s = expand_macros(s)
    s = re.sub(r"(?<!\\)%.*$", "", s, flags=re.M)          # comments
    s = s.replace("\\%", "\uE000").replace("\\&", "\uE001")
    s = s.replace("\\_", "\uE002").replace("\\$", "\uE003")
    s = s.replace("\\#", "\uE004")
    # escaped braces, protected before the brace-stripping pass below
    s = s.replace("\\{", "\uE005").replace("\\}", "\uE006")
    # accents: \"o and friends are not alphabetic commands, so the generic
    # command stripper below would otherwise leave the backslash behind
    for acc, table in (('"', ACCENT_UML), ("'", ACCENT_ACUTE), ("`", ACCENT_GRAVE)):
        s = re.sub(r"\\" + re.escape(acc) + r"\{?([A-Za-z])\}?",
                   lambda m, tb=table: tb.get(m.group(1), m.group(1)), s)

    # maths that survives as text
    s = re.sub(r"\$([^$]*)\$", lambda m: m.group(1), s)
    s = re.sub(r"\^\{([23])\}", lambda m: SUPERS[m.group(1)], s)
    s = re.sub(r"\^([23])\b", lambda m: SUPERS[m.group(1)], s)
    s = re.sub(r"_\{\\mathrm\{([^}]*)\}\}", r"_\1", s)
    s = re.sub(r"_\{([^}]*)\}", r"_\1", s)
    for pat, rep in SYMBOL:
        s = re.sub(pat, rep, s)

    # text-level commands whose content we keep
    for cmd in ("textbf", "textit", "emph", "texttt", "textsc", "mbox", "text"):
        s = re.sub(r"\\" + cmd + r"\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\[A-Za-z]+\{([^{}]*)\}", r"\1", s)        # any other one-arg command
    s = re.sub(r"\\[A-Za-z]+\{\}", "", s)
    s = re.sub(r"\\[A-Za-z]+", "", s)

    s = re.sub(r"\\\\", " ", s)                             # a LaTeX line break
    s = s.replace("``", "\u201c").replace("''", "\u201d")
    s = s.replace("---", "\u2014").replace("--", "\u2013")
    s = s.replace("~", " ").replace("{", "").replace("}", "")
    s = s.replace("\uE000", "%").replace("\uE001", "&")
    s = s.replace("\uE002", "_").replace("\uE003", "$").replace("\uE004", "#")
    s = s.replace("", "{").replace("", "}")
    return re.sub(r"\s+", " ", s).strip()


BOLD_ITALIC = re.compile(r"\\(textbf|textit|emph)\{([^{}]*)\}")


def detex_keep_edges(s):
    """detex, but keep a single space where the fragment had leading or
    trailing whitespace -- otherwise the words either side of an inline
    command are glued together."""
    lead = " " if s[:1].isspace() else ""
    trail = " " if s[-1:].isspace() else ""
    core = detex(s)
    return f"{lead}{core}{trail}" if core else (lead or trail)


def add_rich(par, latex):
    """Add text to a paragraph, keeping \\textbf and \\emph as real runs."""
    latex = expand_macros(latex)
    pos = 0
    for m in BOLD_ITALIC.finditer(latex):
        if m.start() > pos:
            par.add_run(detex_keep_edges(latex[pos:m.start()]))
        run = par.add_run(detex(m.group(2)))
        if m.group(1) == "textbf":
            run.bold = True
        else:
            run.italic = True
        pos = m.end()
    if pos < len(latex):
        par.add_run(detex_keep_edges(latex[pos:]))
    return par


# ------------------------------------------------------------------ parsing
def split_rows(body_tex):
    rows = []
    for line in body_tex.split("\n"):
        line = line.strip()
        if not line or line.startswith("%") or line.startswith("\\addlinespace"):
            continue
        line = re.sub(r"\\\\\s*$", "", line)
        rows.append([detex(c) for c in line.split("&")])
    return rows


def emit_table(doc, block):
    caption = re.search(r"\\caption\{(.*?)\}\s*\n\s*\\label", block, re.S)
    header = None
    for line in block.split("\n"):
        if "&" in line and "\\\\" in line and "\\inputrows" not in line \
                and "multicolumn" not in line and "cmidrule" not in line:
            header = [detex(c) for c in re.sub(r"\\\\.*$", "", line).split("&")]
    src = re.search(r"\\inputrows\{(generated/[A-Za-z0-9_]+)\}", block)
    if not src:
        return
    path = os.path.join(PKG, src.group(1) + ".tex")
    if not os.path.exists(path):
        return
    rows = split_rows(open(path, encoding="utf-8").read())
    if not rows:
        return

    if caption:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        r = p.add_run("Table. ")
        r.bold = True
        add_rich(p, caption.group(1))
        for run in p.runs:
            run.font.size = Pt(9)

    ncol = max(len(r) for r in rows)
    if header and len(header) < ncol:
        header += [""] * (ncol - len(header))
    table = doc.add_table(rows=0, cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if header:
        cells = table.add_row().cells
        for i, txt in enumerate(header[:ncol]):
            cells[i].text = txt
            for par in cells[i].paragraphs:
                for run in par.runs:
                    run.bold = True
                    run.font.size = Pt(8)
    for row in rows:
        cells = table.add_row().cells
        for i, txt in enumerate(row[:ncol]):
            cells[i].text = txt
            for par in cells[i].paragraphs:
                for run in par.runs:
                    run.font.size = Pt(8)


def emit_figure(doc, block):
    img = re.search(r"\\includegraphics\[[^\]]*\]\{([^}]+)\}", block)
    if img:
        path = os.path.join(PKG, img.group(1))
        if os.path.exists(path):
            doc.add_picture(path, width=Cm(16.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = re.search(r"\\caption\{(.*?)\}\s*\n\s*\\label", block, re.S)
    if cap:
        p = doc.add_paragraph()
        r = p.add_run("Figure. ")
        r.bold = True
        add_rich(p, cap.group(1))
        for run in p.runs:
            run.font.size = Pt(9)
        p.paragraph_format.space_after = Pt(10)


def main():
    src = open(os.path.join(PKG, "response.tex"), encoding="utf-8").read()
    body = src[src.index("\\begin{document}") + len("\\begin{document}"):]
    body = body[: body.index("\\end{document}")]

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.2)
        section.top_margin = section.bottom_margin = Cm(2.0)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Response to Reviewers")
    r.bold = True
    r.font.size = Pt(16)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Chemical Language Models for Antioxidant Activity Prediction: "
                    "A Leakage-Controlled Multi-Partition Benchmark with Quantitative "
                    "Molecular Attribution")
    r.italic = True
    r.font.size = Pt(11)

    # walk the body, consuming whole environments where they start
    i = 0
    ENVS = ("comment", "reply", "figure", "table", "itemize")
    while i < len(body):
        env = re.compile(r"\\begin\{(" + "|".join(ENVS) + r")\}").search(body, i)
        cmd = re.compile(r"\\(section|subsection|change|clearpage|maketitle|"
                         r"thispagestyle|noindent|medskip|bigskip|smallskip)\*?").search(body, i)

        nxt = min([m.start() for m in (env, cmd) if m], default=len(body))
        chunk = body[i:nxt].strip()
        if chunk:
            for para in re.split(r"\n\s*\n", chunk):
                text = detex(para)
                if text:
                    add_rich(doc.add_paragraph(), para)
        if nxt >= len(body):
            break

        if env and env.start() == nxt:
            name = env.group(1)
            close = body.index("\\end{" + name + "}", nxt)
            inner = body[env.end():close]
            if name == "comment":
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.7)
                p.paragraph_format.space_before = Pt(9)
                r = p.add_run("Comment.  ")
                r.bold = True
                r.italic = True
                r.font.color.rgb = REPLY_BLUE
                r2 = p.add_run(detex(inner))
                r2.italic = True
                r2.font.color.rgb = REPLY_BLUE
            elif name == "reply":
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                r = p.add_run("Response.  ")
                r.bold = True
                add_rich(p, inner)
            elif name == "itemize":
                for item in re.split(r"\\item\b", inner)[1:]:
                    p = doc.add_paragraph(style="List Bullet")
                    add_rich(p, item)
            elif name == "figure":
                emit_figure(doc, inner)
            elif name == "table":
                emit_table(doc, inner)
            i = close + len("\\end{" + name + "}")
            continue

        # a bare command
        kind = cmd.group(1)
        rest = body[cmd.end():]
        if kind == "clearpage":
            doc.add_page_break()
            i = cmd.end()
        elif kind in ("section", "subsection"):
            # \subsection*{...}: the star is not part of the captured name
            rest = rest.lstrip("*")
            if not re.match(r"\s*\{", rest):
                i = cmd.end()
                continue
            depth, j = 0, rest.index("{")
            for j in range(rest.index("{"), len(rest)):
                depth += (rest[j] == "{") - (rest[j] == "}")
                if depth == 0:
                    break
            heading = detex(rest[rest.index("{") + 1:j])
            p = doc.add_heading(heading, level=1 if kind == "section" else 2)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
            i = cmd.end() + j + 1
        elif kind == "change":
            depth, j = 0, rest.index("{")
            for j in range(rest.index("{"), len(rest)):
                depth += (rest[j] == "{") - (rest[j] == "}")
                if depth == 0:
                    break
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(10)
            r = p.add_run("Changes made.  ")
            r.bold = True
            add_rich(p, rest[rest.index("{") + 1:j])
            i = cmd.end() + j + 1
        else:
            # a no-op command: swallow its braced argument as well, or the
            # argument text ends up in the document
            i = cmd.end()
            m = re.match(r"\s*\{", rest)
            if m:
                depth = 0
                for j in range(rest.index("{"), len(rest)):
                    depth += (rest[j] == "{") - (rest[j] == "}")
                    if depth == 0:
                        i = cmd.end() + j + 1
                        break

    out = os.path.join(PKG, "response.docx")
    doc.save(out)
    print(f"wrote {out} ({os.path.getsize(out) / 1e6:.2f} MB)")


if __name__ == "__main__":
    main()
